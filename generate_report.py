import pandas as pd

df = pd.read_csv('financial_transactions.csv')

# --- Data cleaning --- 
df['Department'] = df['Department'].str.strip().str.title()
df['Category'] = df['Category'].str.strip().str.title()

df['Notes'] = ''
df['Flag'] = ''
condition_1 = (df['Amount'] == 18500) & (df['Category'] == 'Professional Services')
condition_2 = (df['Amount'] == 9200) & (df['Category'] == 'Taxes')
condition_high = (df['Amount'] > 5000) & (df['Type'] == 'Expense')
condition_placeholder = (df['Description'].isnull()) 

df.loc[condition_1, 'Notes'] = 'Documented outlier: large one-time purchase, verified legitimate business expense.'
df.loc[condition_2, 'Notes'] = 'Documented outlier: large one-time tax payment, verified legitimate business expense.'
df.loc[condition_high, 'Flag'] = 'High'
df.loc[condition_placeholder, 'Description'] = 'No description provided'
df['Date'] = pd.to_datetime(df['Date'])

from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows

# --- Clean Data Sheet ---
wb = Workbook()
ws = wb.active
ws.title = "Transactions"

for row in dataframe_to_rows(df, index=False, header=True):
    ws.append(row)

from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
    for cell in row:
        cell.number_format = 'YYYY-MM-DD'

for row in ws.iter_rows(min_row=2, min_col=7, max_col=7):
    for cell in row:
        cell.number_format = '$#,##0.00'


for cell in ws[1]:
    cell.font = Font(bold=True)


widths = {
    'A': 15,  
    'B': 14,  
    'C': 14,  
    'D': 20,  
    'E': 10,  
    'F': 40,  
    'G': 14,  
    'H': 45,
    'I': 10,    
}

for col, width in widths.items():
    ws.column_dimensions[col].width = width

from openpyxl.styles import Alignment


for row in ws.iter_rows(min_row=2, min_col=8, max_col=8):
    for cell in row:
        cell.alignment = Alignment(wrap_text=True, vertical='top')


for row_num in range(2, ws.max_row + 1):
    if ws.cell(row=row_num, column=8).value:
        ws.row_dimensions[row_num].height = 30


total_income = round(df[df['Type'] == 'Income']['Amount'].sum(), 2)
total_expense = round(df[df['Type'] == 'Expense']['Amount'].sum(), 2)
net_profit = round(total_income - total_expense, 2)

department_summary = df.groupby(['Department', 'Type'])['Amount'].sum()

# --- Suammary ---
ws_summary = wb.create_sheet("Summary")

ws_summary['A1'] = "Financial Summary"
ws_summary['A1'].font = Font(bold=True, size=14)

ws_summary['A3'] = "Total Income"
ws_summary['B3'] = total_income
ws_summary['A4'] = "Total Expense"
ws_summary['B4'] = total_expense
ws_summary['A5'] = "Net Profit"
ws_summary['B5'] = net_profit

for cell in ['B3', 'B4', 'B5']:
    ws_summary[cell].number_format = '$#,##0.00'

ws_summary['A7'] = "By Department"
ws_summary['A7'].font = Font(bold=True)
ws_summary['A8'] = "Department"
ws_summary['B8'] = "Type"
ws_summary['C8'] = "Amount"
for cell in ['A8', 'B8', 'C8']:
    ws_summary[cell].font = Font(bold=True)

ws_summary.column_dimensions['A'].width = 15
ws_summary.column_dimensions['B'].width = 18
ws_summary.column_dimensions['C'].width = 15

from openpyxl.chart import BarChart, Reference

# --- Amount by Department and Type Chart ---
chart = BarChart()
chart.type = "col"          
chart.title = "Amount by Department and Type"
chart.y_axis.title = "Amount ($)"
chart.x_axis.title = "Department"

data = Reference(ws_summary, min_col=3, min_row=9, max_row=16)
chart.add_data(data, titles_from_data=False)
ws_summary.column_dimensions['D'].width = 22

for i, ((dept, t), amount) in enumerate(department_summary.items(), start=9):
    ws_summary.cell(row=i, column=1, value=dept)
    ws_summary.cell(row=i, column=2, value=t)
    ws_summary.cell(row=i, column=3, value=amount)
    ws_summary.cell(row=i, column=4, value=f"{dept} - {t}")

categories = Reference(ws_summary, min_col=4, min_row=9, max_row=16)
chart.set_categories(categories)

from openpyxl.chart.label import DataLabelList
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import RichTextProperties, Paragraph, ParagraphProperties, CharacterProperties

chart.dataLabels = DataLabelList()
chart.dataLabels.showVal = True
chart.dataLabels.showSerName = False
chart.dataLabels.showCatName = False
chart.dataLabels.showLegendKey = False
chart.dataLabels.numFmt = '$#,##0'
chart.dataLabels.sourceLinked = False
chart.dataLabels.txPr = RichText(

    bodyPr=RichTextProperties(),
    p=[Paragraph(pPr=ParagraphProperties(
        defRPr=CharacterProperties(sz=1000, b=True, solidFill="1A1A1A")
    ), endParaRPr=CharacterProperties(sz=1000, b=True, solidFill="1A1A1A"))]
)

chart.series[0].dLbls = chart.dataLabels

from openpyxl.chart.marker import DataPoint
from openpyxl.chart.shapes import GraphicalProperties

colors_by_type = [
    "C89494", 
    "C89494",  
    "C89494",  
    "C89494",  
    "C89494",  
    "8FBC8F",  
    "C89494",  
    "8FBC8F",  
]

serie = chart.series[0]  
serie.data_points = [
    DataPoint(idx=i, spPr=GraphicalProperties(solidFill=color))
    for i, color in enumerate(colors_by_type)
]

from openpyxl.drawing.line import LineProperties

chart.y_axis.majorGridlines.spPr = GraphicalProperties(
    ln=LineProperties(solidFill="D9D9D9", w=6350)  
)

chart.width = 22   
chart.height = 10

ws_summary.add_chart(chart, "E2")

wb.save("financial_report.xlsx")

